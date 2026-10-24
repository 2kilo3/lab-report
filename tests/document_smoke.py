#!/usr/bin/env python3
"""Optional synthetic DOCX smoke; requires python-docx and separate page rendering.

This creates no course report. Structural checks do NOT certify page appearance.
"""
import argparse
from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from test_memory_workflow import MemoryWorkflowTests, m


AUTHOR = 'Synthetic Test Author'


def protected(document):
    return [p._p.xml for p in document.paragraphs[:2]] + [document.tables[0]._tbl.xml] + [s._sectPr.xml for s in document.sections]


def build(output):
    output.mkdir(parents=True, exist_ok=True)
    template = output / 'synthetic-template.docx'
    d = Document()
    d.sections[0].page_width, d.sections[0].page_height = Inches(8.5), Inches(11)
    d.styles['Normal'].font.name = 'Arial'
    d.styles['Normal'].font.size = Pt(12)
    title = d.add_paragraph('SYNTHETIC TEMPLATE - FIXED COVER')
    title.runs[0].font.size = Pt(18)
    title.runs[0].bold = True
    d.add_paragraph('Protected author target: ' + AUTHOR)
    table = d.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    table.cell(0, 0).text, table.cell(0, 1).text = 'FIXED SCORE AREA', 'DO NOT EDIT'
    table.cell(1, 0).text, table.cell(1, 1).text = 'Synthetic criterion', 'Reserved'
    d.add_section(WD_SECTION_START.NEW_PAGE)
    d.add_paragraph('ALLOWED BODY AREA')
    d.add_paragraph('Synthetic observation: only caption alignment is changed in this fixture.')
    d.add_paragraph('Figure 1. Synthetic caption')
    d.core_properties.author = AUTHOR
    d.core_properties.last_modified_by = AUTHOR
    d.save(template)
    baseline = protected(Document(template))
    records = []
    scenario = MemoryWorkflowTests()
    scenario.setUp()
    try:
        for mode in ('off', 'on'):
            target = output / ('synthetic-memory-' + mode + '.docx')
            shutil.copyfile(template, target)
            app_id = None
            if mode == 'on':
                ref = scenario.protocol.candidate('synthetic-caption-check', kind='template_risk',
                    sections=dict(zip(m.BODY_FIELDS, ['Synthetic template v1', 'Caption alignment differs',
                        'Center only the caption paragraph; preserve shared styles, cover, score table and sections',
                        'To be checked in this synthetic fixture', 'Synthetic current DOCX and page rendering',
                        'Compare protected XML, metadata and actual rendered pages', 'Other templates need current checks'])))
                scenario.protocol.decide(ref)
                scenario.call('rollback', to='produce', reason='Register actual synthetic DOCX fixture')
                scenario.call('submit_draft', paths=[str(target)], evidence=[])
                state = scenario.call('adopt_lesson', lesson=ref,
                    application={'action': 'Center caption paragraph only', 'affected': [str(target) + ':caption']})
                app_id = state['lessons_used'][-1]['applications'][-1]['application_id']
            current = Document(target)
            caption = next(p for p in current.paragraphs if p.text.startswith('Figure 1.'))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            current.core_properties.author = AUTHOR
            current.core_properties.last_modified_by = AUTHOR
            current.save(target)
            reopened = Document(target)
            assert protected(reopened) == baseline, 'Protected cover/table/sections drifted'
            assert reopened.core_properties.author == AUTHOR
            assert reopened.core_properties.last_modified_by == AUTHOR
            if app_id:
                scenario.call('status')
                scenario.call('record_application', application_id=app_id, result='Actual synthetic DOCX saved and structural/metadata checks passed; visual check pending')
            records.append(dict(mode=mode, file=target.name, protected_xml_unchanged=True,
                                author=reopened.core_properties.author,
                                last_modified_by=reopened.core_properties.last_modified_by,
                                visual_status='REQUIRES SEPARATE RENDER AND PAGE INSPECTION'))
        m.atomic(output / 'structural-checks.json', {'synthetic': True, 'checks': records})
        return records
    finally:
        scenario.doCleanups()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path)
    print(build(parser.parse_args().output.resolve()))
